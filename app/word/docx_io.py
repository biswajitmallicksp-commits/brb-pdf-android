"""Read and write Word (.docx) files for the editable document (offline, python-docx).

Kept when opening and saving:
  text with bold / italic / underline / strike / colour / highlight / font / size,
  paragraph alignment, left and first-line indents, space before/after, line spacing, tab stops,
  bulleted and numbered lists, tables (with merged cells and column widths), pictures,
  page breaks, page size and margins.
Not kept (shown as plain text or left out): headers/footers, footnotes, text boxes, shapes, comments,
tracked changes, fields. Old .doc (Word 97-2003) files cannot be opened - save them as .docx first.
"""
from __future__ import annotations

import io
import logging
from pathlib import Path

from PySide6.QtCore import QByteArray, QSizeF, Qt, QUrl
from PySide6.QtGui import (
    QBrush, QColor, QFont, QImage, QTextBlockFormat, QTextCharFormat, QTextCursor, QTextDocument, QTextFormat,
    QTextImageFormat, QTextLength, QTextListFormat, QTextOption, QTextTableFormat,
)

from app.word.common import A4, PX, base_font, family_of, is_marker, page_setup, set_base_font, set_page_setup

log = logging.getLogger("pdfworkbench.word")

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
WP = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
EMU_PER_PT = 12700

HIGHLIGHT = {"yellow": "#ffff00", "green": "#00ff00", "cyan": "#00ffff", "magenta": "#ff00ff", "blue": "#0000ff",
             "red": "#ff0000", "darkBlue": "#000080", "darkCyan": "#008080", "darkGreen": "#008000",
             "darkMagenta": "#800080", "darkRed": "#800000", "darkYellow": "#808000", "darkGray": "#808080",
             "lightGray": "#c0c0c0", "black": "#000000", "white": "#ffffff"}


_STYLE_PROP = QTextFormat.Property.UserProperty + 2     # block property: Word style name ("Title")


class DocxError(Exception):
    """A message meant for the user."""


# ===================================================================== read
def _style_chain(style):
    while style is not None:
        yield style
        style = style.base_style


def _eff(run, para, attr: str, default=None):
    """Effective run property: the run, then the paragraph style chain, then document defaults."""
    v = getattr(run.font, attr, None)
    if v is not None:
        return v
    for st in _style_chain(para.style):
        v = getattr(st.font, attr, None)
        if v is not None:
            return v
    return default


def _doc_default_size(doc) -> float:
    try:
        sz = doc.styles.element.find(f"{W}docDefaults/{W}rPrDefault/{W}rPr/{W}sz")
        if sz is not None:
            return int(sz.get(f"{W}val")) / 2
    except Exception:
        pass
    try:
        normal = doc.styles["Normal"].font.size
        if normal is not None:
            return normal.pt
    except KeyError:
        pass
    return 11.0


def _para_prop(para, attr: str):
    v = getattr(para.paragraph_format, attr, None)
    if v is not None:
        return v
    for st in _style_chain(para.style):
        v = getattr(st.paragraph_format, attr, None)
        if v is not None:
            return v
    return None


def load_docx(path: str | Path) -> QTextDocument:
    path = Path(path)
    if path.suffix.lower() == ".doc":
        raise DocxError("Old Word files (.doc, Word 97-2003) cannot be opened. Open it in Word and use "
                        "File > Save As > Word Document (.docx), then open the .docx here.")
    try:
        import docx
        d = docx.Document(str(path))
    except Exception as exc:
        raise DocxError(f"{path.name} could not be read as a Word document.\n\n{exc}") from exc
    qd = QTextDocument()
    sec = d.sections[0] if d.sections else None
    w = sec.page_width.pt if sec is not None and sec.page_width else A4[0]
    h = sec.page_height.pt if sec is not None and sec.page_height else A4[1]
    margins = tuple((getattr(sec, k).pt if sec is not None and getattr(sec, k) is not None else 72.0)
                    for k in ("left_margin", "top_margin", "right_margin", "bottom_margin"))
    set_page_setup(qd, w, h, margins)                     # type: ignore[arg-type]
    base_size = _doc_default_size(d)
    set_base_font(qd, "Calibri", base_size)
    reader = _Reader(d, qd, base_size, (w - margins[0] - margins[2]))
    reader.run()
    qd.setModified(False)
    qd.clearUndoRedoStacks()
    return qd


class _Reader:
    def __init__(self, d, qd: QTextDocument, base_size: float, text_width_pt: float):
        self.d, self.qd, self.base = d, qd, base_size
        self.cur = QTextCursor(qd)
        self.first = True
        self.page_break_next = False
        self.lists: dict[tuple, object] = {}
        self.images = 0
        self.text_width = text_width_pt

    def run(self) -> None:
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        self.cur.beginEditBlock()
        body = self.d.element.body
        for child in body.iterchildren():
            if child.tag == f"{W}p":
                self.paragraph(Paragraph(child, self.d), self.cur, top_level=True)
            elif child.tag == f"{W}tbl":
                self.table(Table(child, self.d))
        self.cur.endEditBlock()

    # -------------------------------------------------------- paragraphs
    def _block_format(self, para) -> QTextBlockFormat:
        from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
        bf = QTextBlockFormat()
        al = _para_prop(para, "alignment")
        if al == WD_ALIGN_PARAGRAPH.CENTER:
            bf.setAlignment(Qt.AlignmentFlag.AlignHCenter)
        elif al == WD_ALIGN_PARAGRAPH.RIGHT:
            bf.setAlignment(Qt.AlignmentFlag.AlignRight)
        elif al in (WD_ALIGN_PARAGRAPH.JUSTIFY, getattr(WD_ALIGN_PARAGRAPH, "DISTRIBUTE", -1)):
            bf.setAlignment(Qt.AlignmentFlag.AlignJustify)
        li = _para_prop(para, "left_indent")
        if li is not None:
            bf.setLeftMargin(max(0.0, li.pt * PX))
        ri = _para_prop(para, "right_indent")
        if ri is not None:
            bf.setRightMargin(max(0.0, ri.pt * PX))
        fi = _para_prop(para, "first_line_indent")
        if fi is not None:
            bf.setTextIndent(fi.pt * PX)
        sb = _para_prop(para, "space_before")
        sa = _para_prop(para, "space_after")
        bf.setTopMargin((sb.pt if sb is not None else 0) * PX)
        bf.setBottomMargin((sa.pt if sa is not None else 8) * PX)
        ls = _para_prop(para, "line_spacing")
        rule = _para_prop(para, "line_spacing_rule")
        if isinstance(ls, float):
            bf.setLineHeight(ls * 100, QTextBlockFormat.LineHeightTypes.ProportionalHeight.value)
        elif ls is not None and hasattr(ls, "pt"):
            kind = (QTextBlockFormat.LineHeightTypes.MinimumHeight if rule == WD_LINE_SPACING.AT_LEAST
                    else QTextBlockFormat.LineHeightTypes.FixedHeight)
            bf.setLineHeight(ls.pt * PX, kind.value)
        tabs = []
        try:
            for ts in para.paragraph_format.tab_stops:
                kind = {1: QTextOption.TabType.CenterTab, 2: QTextOption.TabType.RightTab,
                        3: QTextOption.TabType.DelimiterTab}.get(int(ts.alignment), QTextOption.TabType.LeftTab) \
                    if ts.alignment is not None else QTextOption.TabType.LeftTab
                tabs.append(QTextOption.Tab(ts.position.pt * PX, kind))
        except Exception:
            pass
        if tabs:
            bf.setTabPositions(tabs)
        if _para_prop(para, "page_break_before") or self.page_break_next:
            bf.setPageBreakPolicy(QTextFormat.PageBreakFlag.PageBreak_AlwaysBefore)
            self.page_break_next = False
        if _para_prop(para, "keep_together"):
            bf.setNonBreakableLines(True)
        return bf

    def _list_for(self, para):
        """Bulleted / numbered lists (by numbering or by style name)."""
        pPr = para._p.pPr
        num_id, level = None, 0
        if pPr is not None and pPr.numPr is not None:
            try:
                num_id = pPr.numPr.numId.val if pPr.numPr.numId is not None else None
                level = pPr.numPr.ilvl.val if pPr.numPr.ilvl is not None else 0
            except Exception:
                num_id = None
        name = (para.style.name or "").lower() if para.style is not None else ""
        if num_id is None and "list" not in name:
            return None
        if num_id == 0:
            return None
        numbered = "number" in name or self._is_numbered(num_id, level)
        key = (num_id if num_id is not None else name, level)
        lst = self.lists.get(key)
        return key, numbered, level, lst

    def _is_numbered(self, num_id, level) -> bool:
        try:
            numbering = self.d.part.numbering_part.element
            num = numbering.num_having_numId(num_id)
            abstract_id = num.abstractNumId.val
            for an in numbering.findall(f"{W}abstractNum"):
                if int(an.get(f"{W}abstractNumId")) == int(abstract_id):
                    for lvl in an.findall(f"{W}lvl"):
                        if int(lvl.get(f"{W}ilvl")) == int(level):
                            fmt = lvl.find(f"{W}numFmt")
                            return fmt is not None and fmt.get(f"{W}val") not in ("bullet", "none")
        except Exception:
            pass
        return False

    def _char_format(self, run, para) -> QTextCharFormat:
        cf = QTextCharFormat()
        name = _eff(run, para, "name") or "Calibri"
        cf.setFontFamilies([name])
        size = _eff(run, para, "size")
        cf.setFontPointSize(size.pt if size is not None else self.base)
        if _eff(run, para, "bold"):
            cf.setFontWeight(QFont.Weight.Bold)
        if _eff(run, para, "italic"):
            cf.setFontItalic(True)
        if _eff(run, para, "underline"):
            cf.setFontUnderline(True)
        if _eff(run, para, "strike"):
            cf.setFontStrikeOut(True)
        try:
            rgb = run.font.color.rgb if run.font.color is not None and run.font.color.type is not None else None
        except Exception:
            rgb = None
        if rgb is None:
            for st in _style_chain(para.style):
                try:
                    if st.font.color is not None and st.font.color.rgb is not None:
                        rgb = st.font.color.rgb
                        break
                except Exception:
                    pass
        if rgb is not None:
            cf.setForeground(QColor(f"#{rgb}"))
        try:
            hl = run.font.highlight_color
            if hl is not None:
                key = str(hl).split(".")[-1].split(" ")[0].lower()
                for k, v in HIGHLIGHT.items():
                    if k.lower() == key.replace("_", ""):
                        cf.setBackground(QBrush(QColor(v)))
        except Exception:
            pass
        va = _eff(run, para, "superscript")
        if va:
            cf.setVerticalAlignment(QTextCharFormat.VerticalAlignment.AlignSuperScript)
        elif _eff(run, para, "subscript"):
            cf.setVerticalAlignment(QTextCharFormat.VerticalAlignment.AlignSubScript)
        return cf

    def paragraph(self, para, cur: QTextCursor, top_level: bool, in_cell_first: bool = False) -> None:
        p = para._p
        only_break = (top_level and not in_cell_first and p.find(f".//{W}t") is None
                      and p.find(f".//{W}drawing") is None and p.find(f".//{W}br[@{W}type='page']") is not None)
        if only_break:                                   # a paragraph that is just a page break
            self.page_break_next = True
            return
        bf = self._block_format(para)
        sname = (para.style.name or "") if para.style is not None else ""
        if sname == "Title":
            bf.setHeadingLevel(1)
            bf.setProperty(_STYLE_PROP, "Title")
        elif sname.startswith("Heading ") and sname[8:].isdigit():
            bf.setHeadingLevel(int(sname[8:]))
        if in_cell_first:
            cur.setBlockFormat(bf)
        elif self.first and top_level:
            cur.setBlockFormat(bf)
            self.first = False
        else:
            cur.insertBlock(bf, QTextCharFormat())
        info = self._list_for(para)
        if info is not None:
            key, numbered, level, lst = info
            if lst is not None:
                lst.add(cur.block())
            else:
                lf = QTextListFormat()
                lf.setStyle(QTextListFormat.Style.ListDecimal if numbered else
                            [QTextListFormat.Style.ListDisc, QTextListFormat.Style.ListCircle,
                             QTextListFormat.Style.ListSquare][level % 3])
                lf.setIndent(level + 1)
                self.lists[key] = cur.createList(lf)
        # runs, pictures and breaks in document order
        for child in para._p.iterchildren():
            tag = child.tag
            if tag == f"{W}r":
                self._run(child, para, cur)
            elif tag == f"{W}hyperlink":
                for r in child.iterchildren(f"{W}r"):
                    self._run(r, para, cur, link=True)

    def _run(self, r_el, para, cur: QTextCursor, link: bool = False) -> None:
        from docx.text.run import Run
        run = Run(r_el, para)
        cf = self._char_format(run, para)
        if link:
            cf.setFontUnderline(True)
            cf.setForeground(QColor("#0563c1"))
        for el in r_el.iterchildren():
            tag = el.tag
            if tag == f"{W}t":
                cur.insertText(el.text or "", cf)
            elif tag == f"{W}tab":
                cur.insertText("\t", cf)
            elif tag in (f"{W}br", f"{W}cr"):
                kind = el.get(f"{W}type")
                if kind == "page":
                    self.page_break_next = True
                else:
                    cur.insertText(" ", cf)
            elif tag == f"{W}drawing":
                self._picture(el, cur)
            elif tag == f"{W}noBreakHyphen":
                cur.insertText("-", cf)

    def _picture(self, drawing, cur: QTextCursor) -> None:
        blip = drawing.find(f".//{A}blip")
        ext = drawing.find(f".//{WP}extent")
        if blip is None:
            return
        rid = blip.get(f"{R}embed")
        try:
            part = self.d.part.related_parts[rid]
            img = QImage()
            img.loadFromData(QByteArray(part.blob))
            if img.isNull():
                return
        except Exception:
            return
        self.images += 1
        name = f"docximg://{self.images}"
        self.qd.addResource(QTextDocument.ResourceType.ImageResource, QUrl(name), img)
        fmt = QTextImageFormat()
        fmt.setName(name)
        if ext is not None:
            wpt, hpt = int(ext.get("cx")) / EMU_PER_PT, int(ext.get("cy")) / EMU_PER_PT
        else:
            wpt, hpt = img.width() * 0.75, img.height() * 0.75
        if wpt > self.text_width:                          # never wider than the page
            hpt *= self.text_width / wpt
            wpt = self.text_width
        fmt.setWidth(wpt * PX)
        fmt.setHeight(hpt * PX)
        cur.insertImage(fmt)

    # ------------------------------------------------------------ tables
    def table(self, table) -> None:
        from docx.text.paragraph import Paragraph
        rows = table.rows
        if not rows:
            return
        grid = table._tbl.tblGrid
        cols = [int(gc.get(f"{W}w") or 0) / 20 for gc in grid.findall(f"{W}gridCol")] if grid is not None else []
        ncols = max(len(cols), max(len(r._tr.tc_lst) for r in rows))
        if self.first:
            self.first = False
        tf = QTextTableFormat()
        tf.setCellPadding(3)
        tf.setCellSpacing(0)
        tf.setBorder(1)
        tf.setBorderBrush(QBrush(QColor(0, 0, 0)))
        tf.setBorderStyle(QTextFrameFormatBorderStyle.BorderStyle_Solid)
        tf.setBorderCollapse(True)
        if cols and sum(cols) > 0:
            tf.setColumnWidthConstraints([QTextLength(QTextLength.Type.FixedLength, c * PX) for c in cols]
                                         + [QTextLength()] * (ncols - len(cols)))
        qt = self.cur.insertTable(len(rows), ncols, tf)
        # walk the real cells (handles gridSpan and vMerge)
        vmerge_start: dict[int, int] = {}
        for ri, row in enumerate(rows):
            ci = 0
            for tc in row._tr.tc_lst:
                span = 1
                vmerge = None
                tcPr = tc.tcPr
                if tcPr is not None:
                    gs = tcPr.find(f"{W}gridSpan")
                    if gs is not None:
                        span = int(gs.get(f"{W}val") or 1)
                    vm = tcPr.find(f"{W}vMerge")
                    if vm is not None:
                        vmerge = vm.get(f"{W}val") or "continue"
                if ci >= ncols:
                    break
                if vmerge == "continue" and ci in vmerge_start:
                    start = vmerge_start[ci]
                    qt.mergeCells(start, ci, ri - start + 1, span)
                    ci += span
                    continue
                if vmerge == "restart":
                    vmerge_start[ci] = ri
                else:
                    vmerge_start.pop(ci, None)
                if span > 1:
                    qt.mergeCells(ri, ci, 1, min(span, ncols - ci))
                cell = qt.cellAt(ri, ci)
                cc = cell.firstCursorPosition()
                paras = [Paragraph(p, self.d) for p in tc.iterchildren(f"{W}p")]
                for k, p in enumerate(paras):
                    self.paragraph(p, cc, top_level=False, in_cell_first=(k == 0))
                ci += span
        self.cur.movePosition(QTextCursor.MoveOperation.End)
        self.first = True                                 # reuse the empty block Qt keeps after a table


# QTextFrameFormat.BorderStyle lives on QTextFrameFormat
from PySide6.QtGui import QTextFrameFormat as _TFF  # noqa: E402
QTextFrameFormatBorderStyle = _TFF.BorderStyle


# ==================================================================== write
def save_docx(doc: QTextDocument, path: str | Path) -> None:
    """Write the document as a .docx: paragraphs, styles, lists, tables, pictures, page breaks, page setup."""
    from docx import Document
    from docx.shared import Pt

    w, h, margins = page_setup(doc)
    out = Document()
    sec = out.sections[0]
    sec.page_width, sec.page_height = Pt(w), Pt(h)
    sec.left_margin, sec.top_margin, sec.right_margin, sec.bottom_margin = (Pt(m) for m in margins)
    normal = out.styles["Normal"]
    fam, size = base_font(doc)
    normal.font.name = fam
    normal.font.size = Pt(size)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)
    body = out.element.body
    for p in list(body.iterchildren(f"{W}p")):           # the template's empty first paragraph
        body.remove(p)
    writer = _Writer(doc, out)
    writer.frame(doc.rootFrame(), out)
    out.core_properties.author = "BRB PDF"
    out.save(str(path))


class _Writer:
    def __init__(self, doc: QTextDocument, out):
        self.doc, self.out = doc, out
        self.break_next = False
        self.any_para = False

    def frame(self, frame, container) -> None:
        from PySide6.QtGui import QTextTable
        it = frame.begin()
        while not it.atEnd():
            child = it.currentFrame()
            block = it.currentBlock()
            if child is not None:
                if isinstance(child, QTextTable):
                    self.table(child, container)
                else:
                    self.frame(child, container)
            elif block.isValid():
                self.block(block, container)
            it += 1

    def block(self, block, container) -> None:
        if is_marker(block):
            self.break_next = self.any_para
            return
        para = container.add_paragraph()
        self.any_para = True
        self.fill_paragraph(para, block)
        if self.break_next:
            para.paragraph_format.page_break_before = True
            self.break_next = False

    def fill_paragraph(self, para, block) -> None:
        from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
        from docx.shared import Pt
        bf = block.blockFormat()
        pf = para.paragraph_format
        pf.space_before = Pt(bf.topMargin() / PX)
        pf.space_after = Pt(bf.bottomMargin() / PX)
        lst = block.textList()
        indent = bf.leftMargin()
        if lst is not None:
            style = lst.format().style()
            numbered = style in (QTextListFormat.Style.ListDecimal, QTextListFormat.Style.ListLowerAlpha,
                                 QTextListFormat.Style.ListUpperAlpha, QTextListFormat.Style.ListLowerRoman,
                                 QTextListFormat.Style.ListUpperRoman)
            level = max(1, lst.format().indent())
            name = ("List Number" if numbered else "List Bullet") + ("" if level == 1 else f" {min(level, 3)}")
            try:
                para.style = self.out.styles[name]
            except KeyError:
                para.style = self.out.styles["List Number" if numbered else "List Bullet"]
        elif bf.headingLevel() > 0:
            name = "Title" if bf.property(_STYLE_PROP) == "Title" else f"Heading {min(bf.headingLevel(), 9)}"
            try:
                para.style = self.out.styles[name]
            except KeyError:
                pass
            if indent:
                pf.left_indent = Pt(indent / PX)
        elif indent:
            pf.left_indent = Pt(indent / PX)
        if bf.rightMargin():
            pf.right_indent = Pt(bf.rightMargin() / PX)
        if bf.textIndent():
            pf.first_line_indent = Pt(bf.textIndent() / PX)
        if bf.pageBreakPolicy() & QTextFormat.PageBreakFlag.PageBreak_AlwaysBefore:
            pf.page_break_before = True
        lh_type = bf.lineHeightType()
        if lh_type == QTextBlockFormat.LineHeightTypes.ProportionalHeight.value and bf.lineHeight() not in (0, 100):
            pf.line_spacing = bf.lineHeight() / 100
        elif lh_type in (QTextBlockFormat.LineHeightTypes.FixedHeight.value,
                         QTextBlockFormat.LineHeightTypes.MinimumHeight.value) and bf.lineHeight() > 0:
            pf.line_spacing = Pt(bf.lineHeight() / PX)
        al = bf.alignment()
        if al & Qt.AlignmentFlag.AlignHCenter:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif al & Qt.AlignmentFlag.AlignJustify:
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif al & Qt.AlignmentFlag.AlignRight:
            para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for tab in bf.tabPositions():
            kind = {QTextOption.TabType.CenterTab: WD_TAB_ALIGNMENT.CENTER,
                    QTextOption.TabType.RightTab: WD_TAB_ALIGNMENT.RIGHT,
                    QTextOption.TabType.DelimiterTab: WD_TAB_ALIGNMENT.DECIMAL}.get(tab.type, WD_TAB_ALIGNMENT.LEFT)
            pf.tab_stops.add_tab_stop(Pt(tab.position / PX), kind)
        it = block.begin()
        while not it.atEnd():
            frag = it.fragment()
            if frag.isValid():
                self.fragment(para, frag)
            it += 1

    def fragment(self, para, frag) -> None:
        from docx.enum.text import WD_COLOR_INDEX
        from docx.oxml.ns import qn
        from docx.shared import Pt, RGBColor
        cf = frag.charFormat()
        if cf.isImageFormat():
            imf = cf.toImageFormat()
            img = self.doc.resource(QTextDocument.ResourceType.ImageResource, QUrl(imf.name()))
            if isinstance(img, QImage) and not img.isNull():
                from PySide6.QtCore import QBuffer, QIODevice
                buf = QBuffer()
                buf.open(QIODevice.OpenModeFlag.WriteOnly)
                img.save(buf, "PNG")
                width = imf.width() / PX if imf.width() > 0 else img.width() * 0.75
                for _ in range(frag.length()):
                    para.add_run().add_picture(io.BytesIO(bytes(buf.data())), width=Pt(width))
            return
        text = frag.text().replace("￼", "")
        if not text:
            return
        parts = text.split(" ")
        for k, part in enumerate(parts):
            if k:
                para.add_run().add_break()
            if not part:
                continue
            run = para.add_run(part)
            fam = family_of(cf)
            if fam:
                run.font.name = fam
                rpr = run._element.get_or_add_rPr()
                rpr.get_or_add_rFonts().set(qn("w:cs"), fam)
            size = cf.fontPointSize()
            if size > 0:
                run.font.size = Pt(size)
            if cf.fontWeight() >= QFont.Weight.DemiBold.value:
                run.bold = True
            if cf.fontItalic():
                run.italic = True
            if cf.fontUnderline():
                run.underline = True
            if cf.fontStrikeOut():
                run.font.strike = True
            va = cf.verticalAlignment()
            if va == QTextCharFormat.VerticalAlignment.AlignSuperScript:
                run.font.superscript = True
            elif va == QTextCharFormat.VerticalAlignment.AlignSubScript:
                run.font.subscript = True
            if cf.foreground().style() != Qt.BrushStyle.NoBrush:
                col = cf.foreground().color()
                if col.isValid() and col != QColor(0, 0, 0):
                    run.font.color.rgb = RGBColor(col.red(), col.green(), col.blue())
            if cf.background().style() != Qt.BrushStyle.NoBrush:
                bg = cf.background().color().name().lower()
                best = min(HIGHLIGHT.items(), key=lambda kv: _dist(kv[1], bg))[0]
                idx = {"yellow": WD_COLOR_INDEX.YELLOW, "green": WD_COLOR_INDEX.BRIGHT_GREEN,
                       "cyan": WD_COLOR_INDEX.TURQUOISE, "magenta": WD_COLOR_INDEX.PINK,
                       "blue": WD_COLOR_INDEX.BLUE, "red": WD_COLOR_INDEX.RED,
                       "lightGray": WD_COLOR_INDEX.GRAY_25, "darkGray": WD_COLOR_INDEX.GRAY_50}.get(best)
                if idx is not None and bg not in ("#ffffff",):
                    run.font.highlight_color = idx

    def table(self, table, container) -> None:
        from docx.shared import Pt
        rows, cols = table.rows(), table.columns()
        t = container.add_table(rows=rows, cols=cols)
        try:
            t.style = self.out.styles["Table Grid"]
        except KeyError:
            pass
        widths = table.format().columnWidthConstraints()
        done = set()
        for r in range(rows):
            for c in range(cols):
                cell = table.cellAt(r, c)
                key = (cell.row(), cell.column())
                if key in done:
                    continue
                done.add(key)
                rs, cs = cell.rowSpan(), cell.columnSpan()
                target = t.cell(r, c)
                if rs > 1 or cs > 1:
                    target = target.merge(t.cell(r + rs - 1, c + cs - 1))
                # cell paragraphs
                first = True
                it = cell.begin()
                while not it.atEnd():
                    blk = it.currentBlock()
                    if blk.isValid():
                        para = target.paragraphs[0] if first else target.add_paragraph()
                        first = False
                        self.fill_paragraph(para, blk)
                    it += 1
        for c, wc in enumerate(widths[:cols]):
            if wc.type() == QTextLength.Type.FixedLength and wc.rawValue() > 0:
                for cell in t.columns[c].cells:
                    cell.width = Pt(wc.rawValue() / PX)


def _dist(a: str, b: str) -> int:
    ca, cb = QColor(a), QColor(b)
    return abs(ca.red() - cb.red()) + abs(ca.green() - cb.green()) + abs(ca.blue() - cb.blue())


def new_document(size: tuple[float, float] = A4, margins=(72.0, 72.0, 72.0, 72.0)) -> QTextDocument:
    """A blank document (A4, 1-inch margins, Calibri 11)."""
    qd = QTextDocument()
    set_page_setup(qd, size[0], size[1], margins)
    set_base_font(qd, "Calibri", 11)
    qd.setModified(False)
    return qd


_unused = (QSizeF,)
