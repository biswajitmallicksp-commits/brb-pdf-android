"""Build the editable document (QTextDocument) from the page model, and save it
as a Word (.docx) or PDF file.

Units: the page model is in points; QTextDocument layout uses 96-dpi pixels
(PX = 96/72 per point); font sizes stay in points.
"""
from __future__ import annotations

from pathlib import Path

from PySide6.QtCore import QSizeF, Qt
from PySide6.QtGui import (
    QColor, QFont, QTextBlockFormat, QTextCharFormat, QTextCursor,
    QTextDocument, QTextFormat, QTextOption,
)

from app.pdf_editor.doc_model import PageModel

PX = 96 / 72
MARGIN_PT = 36.0
PAGE_MARK = QTextFormat.Property.UserProperty + 1      # block property: page number (0-based) of a page marker
FAMILIES = {"sans": "Arial", "serif": "Times New Roman", "mono": "Courier New"}


def is_marker(block) -> bool:
    return block.blockFormat().hasProperty(PAGE_MARK)


def marker_text(page: int) -> str:
    return f"────  Page {page + 1}  ────"


# ================================================================== build
def build_document(pages: list[PageModel]) -> QTextDocument:
    doc = QTextDocument()
    first = pages[0] if pages else PageModel(0, 595, 842)
    doc.setDocumentMargin(MARGIN_PT * PX)
    doc.setPageSize(QSizeF(first.width * PX, first.height * PX))
    doc.setProperty("page_size_pt", (first.width, first.height))
    font = QFont(FAMILIES["sans"])
    font.setPointSizeF(10)
    doc.setDefaultFont(font)
    cur = QTextCursor(doc)
    cur.beginEditBlock()
    started = False

    def new_block(bf: QTextBlockFormat, cf: QTextCharFormat | None = None) -> None:
        nonlocal started
        if started:
            cur.insertBlock(bf, cf or QTextCharFormat())
        else:
            cur.setBlockFormat(bf)
            if cf is not None:
                cur.setBlockCharFormat(cf)
            started = True

    for k, pm in enumerate(pages):
        # page marker: a grey line in the editor, a page break in Word / PDF
        bf = QTextBlockFormat()
        bf.setProperty(PAGE_MARK, pm.index)
        bf.setAlignment(Qt.AlignmentFlag.AlignHCenter)
        bf.setTopMargin(18 if k else 0)
        bf.setBottomMargin(6)
        if k:
            bf.setPageBreakPolicy(QTextFormat.PageBreakFlag.PageBreak_AlwaysBefore)
        mf = QTextCharFormat()
        mf.setForeground(QColor(150, 150, 150))
        mf.setFontPointSize(8)
        new_block(bf, mf)
        cur.insertText(marker_text(pm.index), mf)
        cur.block().setUserState(pm.index)
        if not pm.has_text:
            bf = QTextBlockFormat()
            cf = QTextCharFormat()
            cf.setForeground(QColor(160, 60, 0))
            new_block(bf, cf)
            cur.insertText(f"[Page {pm.index + 1} has no text. If it is a scanned page, run OCR first.]", cf)
            cur.block().setUserState(pm.index)
            continue
        for p in pm.paras:
            bf = QTextBlockFormat()
            bf.setTopMargin(min(p.space_before, 72) * PX * 0.75)
            bf.setLineHeight(100, QTextBlockFormat.LineHeightTypes.ProportionalHeight.value)
            if p.align == "center":
                bf.setAlignment(Qt.AlignmentFlag.AlignHCenter)
            texts = []
            if p.is_row:
                tabs = []
                for s in p.segs:
                    pos = (s.x - MARGIN_PT) * PX
                    if pos > 2 and (not tabs or pos > tabs[-1].position + 4):
                        tabs.append(QTextOption.Tab(pos, QTextOption.TabType.LeftTab))
                bf.setTabPositions(tabs)
                bf.setNonBreakableLines(True)              # a table row stays on one line
                lead = (p.segs[0].x - MARGIN_PT) * PX > 2
                for j, s in enumerate(p.segs):
                    texts.append(("\t" if (j or lead) else "", s))
            else:
                if p.align != "center":
                    bf.setLeftMargin(max(0.0, (p.x - MARGIN_PT) * PX))
                texts.append(("", p.segs[0]))
            new_block(bf, QTextCharFormat())
            for prefix, s in texts:
                for i, r in enumerate(s.runs):
                    cf = QTextCharFormat()
                    cf.setFontFamilies([FAMILIES.get(r.family, "Arial")])
                    cf.setFontPointSize(max(4.0, r.size))
                    cf.setFontWeight(QFont.Weight.Bold if r.bold else QFont.Weight.Normal)
                    cf.setFontItalic(r.italic)
                    cf.setForeground(QColor.fromRgbF(*r.color))
                    cur.insertText((prefix if i == 0 else "") + r.text, cf)
            cur.block().setUserState(pm.index)
    cur.endEditBlock()
    doc.setModified(False)
    doc.clearUndoRedoStacks()
    return doc


def page_size_pt(doc: QTextDocument) -> tuple[float, float]:
    v = doc.property("page_size_pt")
    if v:
        return float(v[0]), float(v[1])
    s = doc.pageSize()
    return s.width() / PX, s.height() / PX


# =================================================================== PDF
_FAMILY_KEYS = {"arial": "sans", "times new roman": "serif", "courier new": "mono", "calibri": "calibri",
                "nirmala ui": "indic", "helvetica": "sans", "times": "serif", "courier": "mono"}


def _x_of(line, pos: int) -> float:
    v = line.cursorToX(pos)
    return float(v[0] if isinstance(v, tuple) else v)


def save_pdf(doc: QTextDocument, path: str | Path) -> None:
    """Write the document as a PDF with REAL text (searchable and copyable), placed exactly where the
    editor shows it: the editor's own layout gives every line's position; MuPDF writes the text with the
    same fonts. Page markers become page breaks and are not printed."""
    import re

    import pymupdf as fitz

    from app.pdf_editor import content, fonts

    w, h = page_size_pt(doc)
    clone = doc.clone()
    clone.setDocumentMargin(MARGIN_PT * PX)
    clone.setPageSize(QSizeF(w * PX, h * PX))
    # marker lines keep their page break but take no room and print nothing
    block = clone.begin()
    while block.isValid():
        if is_marker(block):
            c = QTextCursor(block)
            c.movePosition(QTextCursor.MoveOperation.StartOfBlock)
            c.movePosition(QTextCursor.MoveOperation.EndOfBlock, QTextCursor.MoveMode.KeepAnchor)
            c.removeSelectedText()
            fmt = block.blockFormat()
            fmt.setTopMargin(0)
            fmt.setBottomMargin(0)
            c.setBlockFormat(fmt)
            small = QTextCharFormat()
            small.setFontPointSize(1)
            c.setBlockCharFormat(small)
        block = block.next()
    # lay out at exactly 96 dpi, whatever the screen's scaling (so sizes and positions are the same everywhere)
    from PySide6.QtGui import QImage
    ref = QImage(8, 8, QImage.Format.Format_RGB32)
    ref.setDotsPerMeterX(round(96 / 0.0254))
    ref.setDotsPerMeterY(round(96 / 0.0254))
    layout = clone.documentLayout()
    layout.setPaintDevice(ref)
    clone.setPageSize(QSizeF(w * PX, h * PX))
    page_h = h * PX
    out = fitz.open()

    def page_no(n: int) -> fitz.Page:
        while out.page_count <= n:
            out.new_page(width=w, height=h)
        return out[n]
    page_no(0)
    used_fonts: dict[int, set] = {}

    def write(pg: fitz.Page, text: str, x: float, y: float, cf: QTextCharFormat) -> None:
        size = cf.fontPointSize() or 10.0
        fams = cf.fontFamilies() or [cf.fontFamily()]
        key = _FAMILY_KEYS.get((fams[0] if fams else "").lower(), "sans")
        bold = cf.fontWeight() >= QFont.Weight.DemiBold.value
        italic = cf.fontItalic()
        col = cf.foreground().color() if cf.foreground().style() != Qt.BrushStyle.NoBrush else QColor(0, 0, 0)
        rgb = (col.redF(), col.greenF(), col.blueF())
        if fonts.needs_indic_font(text):
            # Bengali / Hindi need shaping: let the HTML engine write it at this place
            style = content.TextStyle("indic", size, rgb, bold, italic)
            box = fitz.Rect(x, y - size * 1.05, min(w - 2, x + size * len(text) * 1.2 + 20), y + size * 0.6)
            try:
                content.add_text(pg, box, text, style, single_line=True)
            except Exception:
                pg.insert_text((x, y), text, fontsize=size, color=rgb)
            return
        fam = fonts.family(key) if any(f.key == key for f in fonts.font_families()) else fonts.family("sans")
        file = fam.file_for(bold, italic)
        if file:
            name = f"PW{fam.key}{'B' if bold else ''}{'I' if italic else ''}"
            font = fitz.Font(fontfile=file)
            pg.insert_text((x, y), text, fontsize=size, fontname=name, fontfile=file, color=rgb)
            used_fonts.setdefault(pg.number, set()).add(name)
        else:
            name = content._BASE14[(fam.base14 or "helv", bold, italic)]
            font = fitz.Font(name)
            pg.insert_text((x, y), text, fontsize=size, fontname=name, color=rgb)
        if cf.fontUnderline():
            uw = font.text_length(text, fontsize=size)
            pg.draw_line((x, y + size * 0.12), (x + uw, y + size * 0.12), color=rgb, width=max(0.5, size / 16))

    block = clone.begin()
    while block.isValid():
        if is_marker(block) or not block.text().strip():
            block = block.next()
            continue
        layout.blockBoundingRect(block)            # makes sure the block is laid out
        tl = block.layout()
        btext = block.text()
        frags = []
        it = block.begin()
        while not it.atEnd():
            f = it.fragment()
            if f.isValid():
                frags.append((f.position() - block.position(), f.length(), f.charFormat()))
            it += 1
        tabs = [t.position for t in block.blockFormat().tabPositions()]
        if tabs and "\t" in btext and tl.lineCount():
            # table row: every piece goes to its own tab stop (exactly the original columns)
            line = tl.lineAt(0)
            top = tl.position().y() + line.y()
            pno = int(top // page_h)
            base = (top - pno * page_h + line.ascent()) / PX
            pg = page_no(pno)
            left = MARGIN_PT + block.blockFormat().leftMargin() / PX
            k = 0
            for fs, fl, cf in frags:
                for j, piece in enumerate(btext[fs:fs + fl].split("\t")):
                    if j:
                        k += 1
                    if piece.strip():
                        x = left if k == 0 else MARGIN_PT + tabs[min(k, len(tabs)) - 1] / PX
                        if k > len(tabs):           # more pieces than tab stops: follow the previous piece
                            x = MARGIN_PT + tabs[-1] / PX + 40 * (k - len(tabs))
                        write(pg, piece.strip(), x, base, cf)
            block = block.next()
            continue
        for li in range(tl.lineCount()):
            line = tl.lineAt(li)
            top = tl.position().y() + line.y()
            pno = int(top // page_h)
            base = (top - pno * page_h + line.ascent()) / PX
            pg = page_no(pno)
            ls, le = line.textStart(), line.textStart() + line.textLength()
            for fs, fl, cf in frags:
                a, b = max(fs, ls), min(fs + fl, le)
                if a >= b:
                    continue
                pos = a
                for piece in re.split(r"(\t|\u2028)", btext[a:b]):
                    if piece and piece not in ("\t", "\u2028") and piece.strip():
                        # keep leading spaces' width, write the visible words
                        lead = len(piece) - len(piece.lstrip(" "))
                        x = (tl.position().x() + _x_of(line, pos + lead)) / PX
                        write(pg, piece.strip(" ").replace("\ufffc", ""), x, base, cf)
                    pos += len(piece)
        block = block.next()
    for pno, names in used_fonts.items():
        for name in names:
            content._fix_space_mapping(out[pno], name)
    try:
        out.subset_fonts()
    except Exception:
        pass
    out.set_metadata({"creator": "BRB PDF", "producer": "BRB PDF"})
    out.save(str(path), garbage=3, deflate=True)
    out.close()


# ================================================================== Word
def save_docx(doc: QTextDocument, path: str | Path) -> None:
    """Write the document as a Word file with the same text, sizes, styles, indents and tab stops."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    w, h = page_size_pt(doc)
    out = Document()
    sec = out.sections[0]
    sec.page_width, sec.page_height = Pt(w), Pt(h)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Pt(MARGIN_PT)
    normal = out.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)

    first_para = True
    break_next = False
    block = doc.begin()
    while block.isValid():
        if is_marker(block):
            break_next = not first_para
            block = block.next()
            continue
        bf = block.blockFormat()
        para = out.add_paragraph() if not (first_para and len(out.paragraphs) == 1 and not out.paragraphs[0].text) \
            else out.paragraphs[0]
        first_para = False
        pf = para.paragraph_format
        pf.space_after = Pt(0)
        pf.space_before = Pt(bf.topMargin() / PX)
        if bf.leftMargin():
            pf.left_indent = Pt(bf.leftMargin() / PX)
        if break_next:
            pf.page_break_before = True
            break_next = False
        align = bf.alignment()
        if align & Qt.AlignmentFlag.AlignHCenter:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align & Qt.AlignmentFlag.AlignRight:
            para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        elif align & Qt.AlignmentFlag.AlignJustify:
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for tab in bf.tabPositions():
            pf.tab_stops.add_tab_stop(Pt(tab.position / PX), WD_TAB_ALIGNMENT.LEFT)
        it = block.begin()
        while not it.atEnd():
            frag = it.fragment()
            if frag.isValid():
                cf = frag.charFormat()
                text = frag.text().replace(" ", "\n").replace("￼", "")
                if text:
                    run = para.add_run(text)
                    fams = cf.fontFamilies() or [cf.fontFamily()]
                    fam = (fams[0] if fams else "") or "Arial"
                    run.font.name = fam
                    rpr = run._element.get_or_add_rPr()
                    rpr.get_or_add_rFonts().set(qn("w:cs"), fam)          # Bengali / Hindi use the same font
                    size = cf.fontPointSize()
                    if size > 0:
                        run.font.size = Pt(size)
                    run.bold = cf.fontWeight() >= QFont.Weight.DemiBold.value or None
                    run.italic = cf.fontItalic() or None
                    run.underline = cf.fontUnderline() or None
                    col = cf.foreground().color()
                    if cf.foreground().style() != Qt.BrushStyle.NoBrush and col.isValid() and col != QColor(0, 0, 0):
                        run.font.color.rgb = RGBColor(col.red(), col.green(), col.blue())
            it += 1
        block = block.next()
    out.core_properties.author = "BRB PDF"
    out.save(str(path))


def plain_text(doc: QTextDocument) -> str:
    """The text without page markers (for tests and Copy All)."""
    lines = []
    block = doc.begin()
    while block.isValid():
        if not is_marker(block):
            lines.append(block.text())
        block = block.next()
    return "\n".join(lines)
