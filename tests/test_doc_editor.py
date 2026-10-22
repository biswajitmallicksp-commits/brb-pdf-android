"""Editable (Word-like) document and the 'show OCR result first' workflow.

Run:  .venv\\Scripts\\python tests\\test_doc_editor.py
"""
from __future__ import annotations

import os
import sys
import tempfile
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")


def pump(app, seconds: float) -> None:
    end = time.time() + seconds
    while time.time() < end:
        app.processEvents()
        time.sleep(0.02)


def statement(path: Path, rotation: int = 0) -> None:
    import pymupdf as fitz
    d = fitz.open()
    for k in range(2):
        p = d.new_page(width=595, height=842)
        p.set_rotation(rotation)

        def put(x, y, t, size=10, font="helv"):
            p.insert_text(fitz.Point(x, y) * p.derotation_matrix, t, fontsize=size, fontname=font, rotate=rotation)
        put(210, 60, "STATE BANK OF INDIA", 16, "hebo")
        put(240, 80, "Account Statement", 12)
        put(50, 110, "Dear Customer, this statement shows all transactions in your savings account for the")
        put(50, 123, "period from 01-04-2025 to 30-06-2025. Please check the entries and inform the branch of")
        put(50, 136, "any discrepancy within 15 days.")
        y = 170
        rows = [("Date", "Description", "Debit", "Credit", "Balance"),
                ("01-04-2025", "Opening balance", "", "", "12,500.00"),
                ("02-04-2025", "NEFT ABC Ltd", "", "5,000.00", "17,500.00"),
                ("05-04-2025", "ATM withdrawal", "2,000.00", "", "15,500.00")]
        for r in rows:
            for x, t in zip((50, 130, 300, 380, 460), r):
                if t:
                    put(x, y, t, 10, "hebo" if r[0] == "Date" else "helv")
            y += 16
        put(50, 260, f"Page {k + 1} end")
    d.save(path)


def main() -> int:  # noqa: C901
    import docx
    import pymupdf as fitz
    from PySide6.QtWidgets import QApplication, QFileDialog, QMessageBox

    from app.database.db import AppDatabase
    from app.pdf_editor import doc_export as X
    from app.pdf_editor import doc_model as M
    from app.pdf_ocr.engine import OcrJobResult
    from app.ui import main_window as mw
    from app.ui import ocr_ui

    app = QApplication.instance() or QApplication(sys.argv)
    tmp = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
    folder = Path(tmp.name)
    src = folder / "statement.pdf"
    statement(src)

    # ---------------------------------------------------------------- 1. model
    for rot in (0, 90, 270):
        p = folder / f"rot{rot}.pdf"
        statement(p, rot)
        pm = M.build_pages(fitz.open(p), [0])[0]
        texts = [para.text for para in pm.paras]
        assert texts[0] == "STATE BANK OF INDIA" and (pm.paras[0].align == "center" or rot), texts
        assert texts[2].startswith("Dear Customer") and texts[2].endswith("within 15 days."), texts  # 3 lines joined
        assert texts[3] == "Date\tDescription\tDebit\tCredit\tBalance", texts
        row = pm.paras[5]
        assert [round(s.x) for s in row.segs] == [50, 130, 380, 460], [s.x for s in row.segs]   # empty Debit skipped
        assert pm.paras[3].segs[0].runs[0].bold and not row.segs[0].runs[0].bold

    # ------------------------------------------------------ 2. Word / PDF files
    qd = X.build_document(M.build_pages(fitz.open(src)))
    out_docx, out_pdf = folder / "o.docx", folder / "o.pdf"
    X.save_docx(qd, out_docx)
    X.save_pdf(qd, out_pdf)
    paras = docx.Document(str(out_docx)).paragraphs
    texts = [p.text for p in paras]
    assert texts.count("STATE BANK OF INDIA") == 2 and "Page" not in texts[0], texts[:3]
    assert "\t02-04-2025\tNEFT ABC Ltd\t5,000.00\t17,500.00" in texts, texts
    second = [p for p in paras if p.text == "STATE BANK OF INDIA"][1]
    assert second.paragraph_format.page_break_before, "page 2 starts on a new page in Word"
    nef = next(p for p in paras if "NEFT" in p.text)
    assert [round(t.position.pt) for t in nef.paragraph_format.tab_stops] == [14, 94, 344, 424], \
        [t.position.pt for t in nef.paragraph_format.tab_stops]
    out = fitz.open(out_pdf)
    assert out.page_count == 2, out.page_count
    t0 = out[0].get_text()
    t1 = out[1].get_text()
    assert "STATE BANK OF INDIA" in t0 and "Page 2 end" not in t0 and "Page 2 end" in t1, (repr(t0[:300]), repr(t1[:300]))
    assert "─" not in t0, "page markers are not printed"
    credit = out[0].search_for("5,000.00")[0]
    assert abs(credit.x0 - 380) < 6, credit                     # the column stays where it was

    # ------------------------------------------------------- 3. the editor window
    QMessageBox.warning = staticmethod(lambda *a, **k: (_ for _ in ()).throw(AssertionError(a)))
    db = AppDatabase(folder / "ui.sqlite3")
    win = mw.MainWindow(db)
    win.resize(1300, 850)
    win.show()
    tab = win.open_file(str(src))
    pump(app, 0.6)
    tab.scanned_pages = []
    win.actions["edit_doc"].trigger()
    pump(app, 0.5)
    ed = win._doc_windows[-1]
    assert ed.isVisible() and "temporary document" in ed.windowTitle(), ed.windowTitle()
    assert "NEFT ABC Ltd" in X.plain_text(ed.doc)
    n = ed.replace_all("NEFT ABC Ltd", "RTGS XYZ Pvt Ltd")
    assert n == 2 and ed.doc.isModified() and ed.windowTitle().startswith("* ")
    saved_word = folder / "edited.docx"
    saved_pdf = folder / "edited.pdf"
    QFileDialog.getSaveFileName = staticmethod(lambda *a, **k: (str(saved_word), ""))
    assert ed.save_word()
    QFileDialog.getSaveFileName = staticmethod(lambda *a, **k: (str(saved_pdf), ""))
    assert ed.save_pdf()
    assert not ed.doc.isModified()
    assert "RTGS XYZ Pvt Ltd" in "\n".join(p.text for p in docx.Document(str(saved_word)).paragraphs)
    assert "RTGS XYZ Pvt Ltd" in fitz.open(saved_pdf)[0].get_text()
    assert "NEFT" in fitz.open(src)[0].get_text(), "the original PDF is never changed"
    ed.grab().save(str(ROOT / "tests" / "samples" / "screenshots" / "doc_editor.png")) \
        if (ROOT / "tests" / "samples" / "screenshots").exists() else None
    # closing with unsaved changes asks; 'Close without saving' closes
    ed.replace_all("Opening balance", "Opening bal.")
    asked = []

    def fake_exec(box):
        asked.append(box.text())
        box._clicked = [b for b in box.buttons() if b.text() == "Close without saving"][0]
        return 0
    QMessageBox.exec = fake_exec
    QMessageBox.clickedButton = lambda box: box._clicked
    ed.close()
    pump(app, 0.2)
    assert asked and ed not in win._doc_windows
    del QMessageBox.exec
    del QMessageBox.clickedButton

    # ----------------------------------- 4. OCR result shown first, saved only when asked
    temp_pdf = ocr_ui.temp_output_path(src)
    fitz.open(src).save(temp_pdf)
    result = OcrJobResult(temp_pdf, pages_done=[0], words=42)
    QMessageBox.exec = lambda box: 0
    win._ocr_finished(result, True, temporary=True, edit_after=True, source=tab)
    del QMessageBox.exec
    pump(app, 0.5)
    t2 = win.current_tab()
    assert t2.doc.temporary and "(not saved)" in win.tabs.tabText(win.tabs.currentIndex()), \
        win.tabs.tabText(win.tabs.currentIndex())
    assert t2.doc.suggested_path.name == "statement_OCR.pdf" and t2.doc.suggested_path.parent == folder
    assert win._doc_windows[-1].isVisible(), "the editable document opens too"
    win._doc_windows[-1].doc.setModified(False)
    win._doc_windows[-1].close()
    # Save on an unsaved OCR result: choose format, folder and the original's fate
    from app.ui import save_options as SO
    seen = {}

    def choose(dlg, fmt="both", delete=False):
        seen["original"] = dlg.original
        seen["name"] = dlg.name_edit.text()
        {"both": dlg.rb_both, "pdf": dlg.rb_pdf, "word": dlg.rb_word}[fmt].setChecked(True)
        dlg.name_edit.setText("kept")
        (dlg.delete if delete else dlg.keep).setChecked(True)
        dlg.choice = SO.SaveChoice(dlg._targets(), delete)
        return 1
    SO.SaveOptionsDialog.exec = lambda dlg: choose(dlg)
    win.save()
    assert seen["original"] == src and seen["name"] == "statement_OCR", seen
    assert (folder / "kept.pdf").exists() and (folder / "kept.docx").exists()
    assert not t2.doc.temporary and src.exists(), "the original is kept unless asked"
    assert "(not saved)" not in win.tabs.tabText(win.tabs.currentIndex())
    win.close_tab(win.tabs.currentIndex())
    pump(app, 0.2)
    assert not temp_pdf.exists(), "the temporary OCR file is removed when its tab closes"


    # editor: one Save window; Word into the same folder, original moved to the Recycle Bin
    victim = folder / "old_statement.pdf"
    statement(victim)
    vt = win.open_file(str(victim))
    pump(app, 0.4)
    vt.scanned_pages = []
    win.edit_as_document(vt)
    pump(app, 0.4)
    ed2 = win._doc_windows[-1]
    assert ed2.original == victim
    ed2.replace_all("ATM withdrawal", "Cash withdrawal")
    assert ed2.a_save.text() == "Save changes..."
    SO.SaveOptionsDialog.exec = lambda dlg: choose(dlg, "word", delete=True)
    QMessageBox.information = staticmethod(lambda *a, **k: 0)
    trashed = []
    SO.move_to_recycle_bin = lambda p: (trashed.append(p) or p.unlink() or (True, ""))
    assert ed2.save()
    assert trashed == [victim] and not victim.exists()
    assert all(t.doc.path != victim for t in (win.tabs.widget(i) for i in range(win.tabs.count())))
    assert "Cash withdrawal" in "\n".join(p.text for p in docx.Document(str(folder / "kept.docx")).paragraphs)
    ed2.close()

    win.close()
    pump(app, 0.2)
    print("Editable document + OCR result test: OK")
    return 0


if __name__ == "__main__":
    sys.exit(main())
