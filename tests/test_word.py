"""Word documents inside BRB PDF: open / edit / save .docx, new documents, PDF export, start screen.

Run:  .venv\\Scripts\\python tests\\test_word.py
"""
from __future__ import annotations

import os
import sys
import tempfile
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")


def pump(app, seconds: float) -> None:
    end = time.time() + seconds
    while time.time() < end:
        app.processEvents()
        time.sleep(0.02)


def sample_docx(path: Path, picture: Path) -> None:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor
    d = docx.Document()
    d.add_heading("Loan Proposal", 0)
    p = d.add_paragraph("Applicant: ")
    r = p.add_run("Ramesh Traders")
    r.bold = True
    r.font.color.rgb = RGBColor(0xC0, 0, 0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.add_paragraph("Term loan of Rs 25 lakh", style="List Bullet")
    d.add_paragraph("Working capital of Rs 10 lakh", style="List Bullet")
    d.add_paragraph("Submit the bank statements", style="List Number")
    t = d.add_table(rows=3, cols=3)
    t.style = "Table Grid"
    for i in range(3):
        for j in range(3):
            t.cell(i, j).text = f"r{i}c{j}"
    d.add_picture(str(picture), width=Pt(90))
    d.add_page_break()
    d.add_paragraph("Second page text")
    d.save(path)


def main() -> int:  # noqa: C901
    import docx
    import pymupdf as fitz
    from PySide6.QtGui import QColor, QImage, QTextCursor
    from PySide6.QtWidgets import QApplication, QFileDialog, QMessageBox

    from app.database.db import AppDatabase
    from app.ui import main_window as mw
    from app.ui.doc_editor import PDF_FILTER, WORD_FILTER
    from app.word import docx_io, pdf_out

    app = QApplication.instance() or QApplication(sys.argv)
    tmp = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
    folder = Path(tmp.name)
    pic = folder / "blue.png"
    img = QImage(60, 30, QImage.Format.Format_RGB32)
    img.fill(QColor("blue"))
    img.save(str(pic))
    src = folder / "proposal.docx"
    sample_docx(src, pic)

    # ------------------------------------------------------------ 1. read / write
    qd = docx_io.load_docx(src)
    text = qd.toPlainText()
    for want in ("Loan Proposal", "Ramesh Traders", "Term loan", "r2c2", "Second page text"):
        assert want in text, (want, text)
    out = folder / "roundtrip.docx"
    docx_io.save_docx(qd, out)
    d = docx.Document(str(out))
    styles = {p.text: p.style.name for p in d.paragraphs}
    assert styles["Loan Proposal"] == "Title", styles
    assert styles["Term loan of Rs 25 lakh"] == "List Bullet", styles
    assert styles["Submit the bank statements"] == "List Number", styles
    ap = next(p for p in d.paragraphs if "Applicant" in p.text)
    bold = [r for r in ap.runs if r.text == "Ramesh Traders"][0]
    assert bold.bold and str(bold.font.color.rgb) == "C00000" and ap.alignment == 1, (bold.bold, ap.alignment)
    assert [[c.text for c in r.cells] for r in d.tables[0].rows][2] == ["r2c0", "r2c1", "r2c2"]
    assert len(d.inline_shapes) == 1 and abs(d.inline_shapes[0].width.pt - 90) < 1
    assert next(p for p in d.paragraphs if p.text == "Second page text").paragraph_format.page_break_before
    pdf = folder / "proposal.pdf"
    pdf_out.save_pdf(qd, pdf)
    pd = fitz.open(pdf)
    assert pd.page_count == 2, pd.page_count
    t0 = pd[0].get_text()
    for want in ("Ramesh Traders", "\u2022", "r1c1"):
        assert want in t0, (want, ascii(t0))
    assert "Second page" not in t0, ascii(t0)
    assert "Second page text" in pd[1].get_text()
    assert pd[0].get_images(), "the picture is in the PDF"
    assert len(pd[0].get_drawings()) >= 4, "table borders are drawn"
    # old Word format: a clear message
    try:
        docx_io.load_docx(folder / "old.doc")
        raise AssertionError("a .doc must be refused")
    except docx_io.DocxError as exc:
        assert "Save As" in str(exc)

    # ------------------------------------------------------------ 2. main window
    QMessageBox.warning = staticmethod(lambda *a, **k: (_ for _ in ()).throw(AssertionError(a)))
    db = AppDatabase(folder / "ui.sqlite3")
    win = mw.MainWindow(db)
    win.resize(1300, 850)
    win.show()
    pump(app, 0.3)
    assert win.stack.currentIndex() == 0 and win.start_page.isVisible(), "start screen when nothing is open"
    # open a .docx the same way as a PDF (File > Open, double-click, command line)
    assert win.open_file(str(src)) is None
    pump(app, 0.3)
    ed = win._doc_windows[-1]
    assert ed.word_mode and ed.path.resolve() == src.resolve() and ed.isVisible(), ed.windowTitle()
    assert ed.windowTitle() == "proposal.docx - BRB PDF Word", ed.windowTitle()
    assert win.open_word(str(src)) is ed, "opening it again brings the same window"
    assert str(src.resolve()) in db.recent_files(), db.recent_files()
    # edit: bold text, a list, a table with an extra row, a picture, a page break
    cur = ed.edit.textCursor()
    cur.movePosition(QTextCursor.MoveOperation.End)
    ed.edit.setTextCursor(cur)
    ed.insert_page_break()
    ed.edit.insertPlainText("Added by test")
    ed.edit.moveCursor(QTextCursor.MoveOperation.StartOfBlock)
    ed.edit.moveCursor(QTextCursor.MoveOperation.EndOfBlock, QTextCursor.MoveMode.KeepAnchor)
    ed.a_bold.trigger()
    ed.a_center.trigger()
    ed.edit.moveCursor(QTextCursor.MoveOperation.End)
    ed.edit.insertPlainText("\nChecklist item")
    ed.a_bullets.trigger()
    ed.edit.moveCursor(QTextCursor.MoveOperation.End)
    ed.edit.textCursor().insertBlock()
    ed.a_bullets.trigger()                   # second press: out of the list again
    table = ed.insert_table(2, 3)
    table.cellAt(0, 0).firstCursorPosition().insertText("Bank")
    c = table.cellAt(1, 2).firstCursorPosition()
    ed.edit.setTextCursor(c)
    ed._table_op("row_below")
    assert table.rows() == 3
    ed.edit.moveCursor(QTextCursor.MoveOperation.End)
    ed.insert_image(img)
    pic_cursor = ed._image_at_cursor()
    assert pic_cursor is not None
    ed.resize_picture(pic_cursor, 2.0)
    ed.set_page(842, 595, (50, 50, 50, 50))     # A4 landscape
    assert ed.doc.isModified() and ed.windowTitle().startswith("* ")
    pump(app, 1.0)
    assert ed.page_count >= 3, ed.page_count
    assert ed.save(), "Ctrl+S writes the same .docx"
    assert not ed.doc.isModified()
    d = docx.Document(str(src))
    added = next(p for p in d.paragraphs if p.text == "Added by test")
    assert added.runs[0].bold and added.alignment == 1 and added.paragraph_format.page_break_before
    assert next(p for p in d.paragraphs if p.text == "Checklist item").style.name == "List Bullet"
    assert len(d.tables) == 2 and len(d.tables[1].rows) == 3 and d.tables[1].cell(0, 0).text == "Bank"
    assert len(d.inline_shapes) == 2
    assert abs(d.sections[0].page_width.pt - 842) < 1 and abs(d.sections[0].left_margin.pt - 50) < 1
    # reopen: everything still there
    again = docx_io.load_docx(src)
    assert "Added by test" in again.toPlainText() and "Bank" in again.toPlainText()
    # Save As PDF (an export: the Word file stays the document)
    target = folder / "export.pdf"
    QFileDialog.getSaveFileName = staticmethod(lambda *a, **k: (str(target), PDF_FILTER))
    ed.edit.insertPlainText("x")
    assert ed.save_as() and target.exists() and ed.doc.isModified() and ed.path.resolve() == src.resolve()
    ed.edit.undo()
    ed.doc.setModified(False)
    ed.grab().save(str(ROOT / "tests" / "samples" / "screenshots" / "word_editor.png")) \
        if (ROOT / "tests" / "samples" / "screenshots").exists() else None
    ed.close()
    pump(app, 0.2)
    assert ed not in win._doc_windows

    # ------------------------------------------------------------ 3. new document
    new = win.new_word_document()
    assert new.word_mode and new.path is None
    new.edit.insertPlainText("Fresh letter to the bank manager")
    new.apply_style("Heading 1")
    saved = folder / "letter.docx"
    QFileDialog.getSaveFileName = staticmethod(lambda *a, **k: (str(saved.with_suffix("")), WORD_FILTER))
    assert new.save(), "first Ctrl+S of a new document asks where (Save As)"
    assert new.path == saved and saved.exists() and new.windowTitle() == "letter.docx - BRB PDF Word"
    h = docx.Document(str(saved)).paragraphs[0]
    assert h.style.name == "Heading 1" and h.text.startswith("Fresh letter"), h.style.name
    # closing with unsaved changes asks: Don't Save closes
    new.edit.insertPlainText(" more")
    asked = []

    def fake_exec(box):
        asked.append(box.text())
        box._clicked = [b for b in box.buttons() if b.text() == "Don't Save"][0]
        return 0
    QMessageBox.exec = fake_exec
    QMessageBox.clickedButton = lambda box: box._clicked
    new.close()
    pump(app, 0.2)
    assert asked and "letter.docx" in asked[0]
    del QMessageBox.exec
    del QMessageBox.clickedButton

    # ------------------------------------------------------------ 4. convert PDF to Word
    pdf_src = folder / "statement.pdf"
    p = fitz.open()
    pg = p.new_page()
    pg.insert_text((60, 80), "Statement of account 01-04-2025", fontsize=12)
    p.save(pdf_src)
    QFileDialog.getOpenFileName = staticmethod(lambda *a, **k: (str(pdf_src), ""))
    win.convert_pdf_to_word()
    pump(app, 0.4)
    conv = win._doc_windows[-1]
    assert not conv.word_mode and "Statement of account" in conv.doc.toPlainText()
    conv.close()
    win.close()
    pump(app, 0.2)
    print("Word documents test: OK")
    return 0


if __name__ == "__main__":
    sys.exit(main())
